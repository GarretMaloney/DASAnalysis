# -*- coding: utf-8 -*-
"""
AnalyzeStrainSaveFigs.py
========================
Loads processed .npz file from ProcessTDMS.py and performs:
  - Strain rate / strain waterfall plots
  - FFT amplitude spectrum (one subplot per section)
  - Displacement per cable section
  - Peak envelope analysis and strain transfer ratio (unfiltered + bandpass)
  - Channel amplitude profile along fiber

Each section measurement averages ch_avg channels centered on the specified channel.
Section boundaries are drawn on waterfall and amplitude profile figures.
Number of sections is fully dynamic — works with any number of inner_channels.

Figures are saved to a 'figures' folder alongside 'processed'.
Results are appended to a shared Excel file for cross-experiment comparison.

Edit the CONFIGURATION block for each run. No processing is redone.
Memory handling: mmap_mode load, float32 throughout, waterfall and Fig 6
downsampled to avoid OOM on large overnight datasets.
Uses sosfiltfilt throughout for numerical stability at very low frequencies.
"""

import os
import re
import glob
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import MultipleLocator
from scipy.signal import detrend, butter, sosfiltfilt, find_peaks
from scipy.interpolate import interp1d

# ============================================================
# CONFIGURATION — edit these between runs
# ============================================================
processed_directory = r"D:\Garret's Experiments\1x Strain Gradient\FLUTe 2.5PSI\0.001Hz\7V\processed"

excel_path = r"D:\Garret's Experiments\1x Strain Gradient\FLUTe 2.5PSI\StrainTransferResults.xlsx"

Freq = 0.001 #None   # float to override, None = auto-detect

# Section channel midpoints, cable names, and colors
# Add or remove entries as needed — everything scales automatically
sections = {
    100: ('Reference', 'black'),
    195: ('Yellow',    'gold'),
    300: ('Green',     'tab:green'),
    370: ('Blue',      'tab:blue'),
    425: ('Steel',     'tab:red'),
}
ref_channel    = 100
inner_channels = [195, 300, 370, 425]

def section_name(ch):
    return sections[ch][0]

def section_color(ch):
    return sections[ch][1]

def section_label(ch):
    return f'{section_name(ch)} (Ch {ch})'

# Number of channels to average for each section measurement
ch_avg = 10

# Bandpass
bp_low   = None
bp_high  = None
bp_order = 4

# Max samples for waterfall display (downsamples in time only)
waterfall_max_samples = 50000

# Max samples for time-series plots (Figs 4 & 5) — downsamples for rendering
plot_max_samples = 50000

# Downsample factor for Fig 6 per-channel loop (reduces filter cost)
fig6_ds = 10

# Set to a depth in meters to crop fiber beyond that point, None = use all channels
max_depth_m = 115.0
# ============================================================


# ============================================================
# UTILITIES
# ============================================================
def cumtrapz(y, dx=1.0, axis=-1):
    y = np.moveaxis(y, axis, -1)
    out = np.cumsum((y[..., 1:] + y[..., :-1]) * 0.5, axis=-1) * dx
    out = np.concatenate([np.zeros((*out.shape[:-1], 1), dtype=out.dtype), out], axis=-1)
    return np.moveaxis(out, -1, axis)

def bandpass(data, fs, lowcut, highcut, order=2):
    nyq = 0.5 * fs
    sos = butter(order, [lowcut / nyq, highcut / nyq], btype='band', output='sos')
    return sosfiltfilt(sos, data, axis=-1)

def nextpow2(n):
    return int(np.ceil(np.log2(max(n, 1))))

def peak_envelope(x, fs, freq_hz):
    tloc     = np.arange(x.size) / fs
    min_dist = max(1, int(round(0.8 * fs / max(freq_hz, 1e-9))))
    p_idx,  _ = find_peaks( x, distance=min_dist)
    tr_idx, _ = find_peaks(-x, distance=min_dist)
    if p_idx.size >= 2:
        upper = interp1d(tloc[p_idx],  x[p_idx],  kind='linear',
                         fill_value='extrapolate', bounds_error=False)(tloc)
    else:
        upper = np.maximum.accumulate(x)
    if tr_idx.size >= 2:
        lower = interp1d(tloc[tr_idx], x[tr_idx], kind='linear',
                         fill_value='extrapolate', bounds_error=False)(tloc)
    else:
        lower = -np.maximum.accumulate(-x)
    return upper, lower, (upper - lower)

def savefig(fig, figures_dir, name):
    path = os.path.join(figures_dir, name)
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)
    print(f"  Saved: {path}")

def parse_hz(s):
    m = re.search(r'(\d+\.?\d*)\s*hz', s.lower())
    return float(m.group(1)) if m else None

def parse_voltage(s):
    m = re.search(r'(\d+\.?\d*)\s*v', s.lower())
    return float(m.group(1)) if m else None


# ============================================================
# SETUP DIRECTORIES
# ============================================================
base_dir    = os.path.dirname(processed_directory.rstrip(os.sep))
figures_dir = os.path.join(base_dir, 'figures')
os.makedirs(figures_dir, exist_ok=True)

path_parts  = processed_directory.replace('\\', '/').split('/')
vol_folder  = path_parts[-2] if path_parts[-1] == 'processed' else path_parts[-1]
freq_folder = path_parts[-3] if path_parts[-1] == 'processed' else path_parts[-2]
freq_val    = parse_hz(freq_folder)
voltage_val = parse_voltage(vol_folder)

# ============================================================
# LOAD DATA — tries decdata.npy (true mmap) first,
# falls back to .npz for legacy datasets
# ============================================================
npy_file  = os.path.join(processed_directory, 'decdata.npy')
meta_file = os.path.join(processed_directory, 'metadata.npz')

if os.path.exists(npy_file) and os.path.exists(meta_file):
    print(f"Loading: decdata.npy from {processed_directory}")
    decdata     = np.load(npy_file, mmap_mode='r')
    meta        = np.load(meta_file)
    fs_d        = float(meta['fs_d'])
    spatial_res = float(meta['spatial_samp']) if 'spatial_samp' in meta.files else 0.25
    load_label  = 'decdata.npy'
else:
    npz_files = sorted(glob.glob(os.path.join(processed_directory, '*.npz')))
    if not npz_files:
        raise FileNotFoundError(f"No decdata.npy or .npz found in {processed_directory}")
    print(f"Loading: {os.path.basename(npz_files[0])} from {processed_directory} (legacy npz)")
    npz         = np.load(npz_files[0])
    decdata     = npz['decdata']
    fs_d        = float(npz['fs_d'])
    spatial_res = float(npz['spatial_samp']) if 'spatial_samp' in npz.files else 0.25
    load_label  = os.path.basename(npz_files[0])

n_channels = decdata.shape[1]
print(f"  Shape: {decdata.shape}  |  fs_d: {fs_d} Hz  |  Duration: {decdata.shape[0]/fs_d:.1f} s")
print(f"  Channels: {n_channels}  |  spatial_res: {spatial_res} m")

# ============================================================
# SETUP — crop depth, transpose to [ch x time], keep float32
# ============================================================
if max_depth_m is not None:
    ch_stop = int(np.floor(max_depth_m / spatial_res))
    ch_stop = min(ch_stop, decdata.shape[1])
else:
    ch_stop = decdata.shape[1]

subdata    = decdata[:, :ch_stop].T.astype(np.float32)
nch        = subdata.shape[0]
L          = subdata.shape[1]
t          = np.arange(L) / fs_d
depth_axis = np.arange(ch_stop) * spatial_res

all_channels = [ref_channel] + inner_channels
for ch in all_channels:
    if ch >= nch:
        raise ValueError(
            f"Channel {ch} ({section_name(ch)}) out of range — "
            f"data only has {nch} channels (0–{nch-1})"
        )

# Plot downsampling step for Figs 4 & 5 (rendering only, not analysis)
plot_step = max(1, int(np.ceil(L / plot_max_samples)))
t_plot    = t[::plot_step]
if plot_step > 1:
    print(f"  Plot downsampled {plot_step}x → {len(t_plot)} samples for Figs 4/5")

# ============================================================
# SECTION LINE HELPER
# ============================================================
def draw_section_lines(ax, orientation='depth'):
    half = ch_avg // 2
    for ch in all_channels:
        col = section_color(ch)
        lo  = max(0, ch - half)
        hi  = min(nch - 1, ch + half)
        if orientation == 'depth':
            ax.axhline(lo * spatial_res, color=col, lw=0.8, linestyle='--', alpha=0.8)
            ax.axhline(hi * spatial_res, color=col, lw=0.8, linestyle='--', alpha=0.8)
            ax.axhline(ch * spatial_res, color=col, lw=1.2, linestyle='-',  alpha=0.9,
                       label=section_label(ch))
        else:
            ax.axvline(lo, color=col, lw=0.8, linestyle='--', alpha=0.8)
            ax.axvline(hi, color=col, lw=0.8, linestyle='--', alpha=0.8)
            ax.axvline(ch, color=col, lw=1.2, linestyle='-',  alpha=0.9,
                       label=section_label(ch))

# ============================================================
# AUTO-DETECT FREQUENCY
# ============================================================
half         = ch_avg // 2
ref_lo       = max(0, ref_channel - half)
ref_hi       = min(nch, ref_channel + half)
RateOut_mean = subdata[ref_lo:ref_hi, :].mean(axis=0).astype(np.float64)
RateOut_mean -= RateOut_mean.mean()
DispOut_raw  = cumtrapz(RateOut_mean, dx=1/fs_d) * 1e-3

NFFT   = 2 ** (max(nextpow2(L + 1) - 1, 1))
f_axis = fs_d / 2 * np.linspace(0, 1, NFFT // 2 + 1)
Amp_detect = 2 * np.abs(np.fft.fft(detrend(DispOut_raw, type='linear'), n=NFFT)[:NFFT//2+1]) / NFFT

if Freq is None:
    min_freq_idx = max(1, int(0.0005 * NFFT / fs_d))
    detected_idx = min_freq_idx + int(np.argmax(Amp_detect[min_freq_idx:]))
    Freq = float(f_axis[detected_idx])
    print(f"  Auto-detected frequency : {Freq:.4f} Hz")
else:
    print(f"  Using specified frequency: {Freq:.4f} Hz")

if bp_low  is None: bp_low  = max(Freq * 0.5, 0.0001)
if bp_high is None: bp_high = min(Freq * 1.5, fs_d / 2 * 0.99)
print(f"  Bandpass: {bp_low:.4f} – {bp_high:.4f} Hz")
print(f"  Channel averaging: {ch_avg} channels per section")

# ============================================================
# WATERFALL — downsample in time to cap memory/render cost
# ============================================================
wf_step = max(1, int(np.ceil(L / waterfall_max_samples)))
wf_data = subdata[:, ::wf_step]
wf_t    = t[::wf_step]
wf_fs   = fs_d / wf_step
if wf_step > 1:
    print(f"  Waterfall downsampled {wf_step}x → {wf_data.shape[1]} samples")

wf_rate = wf_data.astype(np.float64)
wf_disp = cumtrapz(wf_rate, dx=wf_step/fs_d, axis=1) * 1e-3
wf_filt = bandpass(wf_disp, wf_fs, bp_low, bp_high, bp_order).astype(np.float32)
del wf_rate, wf_disp

# ============================================================
# FIGURE 1 — Strain Rate Waterfall
# ============================================================
fig1, ax1 = plt.subplots(figsize=(12, 6))
im1 = ax1.imshow(wf_data, aspect='auto', origin='upper', cmap='jet',
                 extent=[wf_t[0], wf_t[-1], depth_axis[-1], depth_axis[0]],
                 vmin=np.nanpercentile(wf_data, 20),
                 vmax=np.nanpercentile(wf_data, 80))
plt.colorbar(im1, ax=ax1, label='Strain Rate (nm/sec)')
ax1.set_xlabel('Time (s)'); ax1.set_ylabel('Depth (m)')
ax1.set_title(f'Strain Rate Waterfall — {ch_avg}-ch section windows shown')
draw_section_lines(ax1, orientation='depth')
ax1.legend(fontsize=7, loc='upper right')
plt.tight_layout()
savefig(fig1, figures_dir, 'fig1_strain_rate_waterfall.png')

# ============================================================
# FIGURE 2 — Filtered Strain Waterfall
# ============================================================
fig2, ax2 = plt.subplots(figsize=(12, 6))
im2 = ax2.imshow(wf_filt, aspect='auto', origin='upper', cmap='jet',
                 extent=[wf_t[0], wf_t[-1], depth_axis[-1], depth_axis[0]],
                 vmin=np.nanpercentile(wf_filt, 20),
                 vmax=np.nanpercentile(wf_filt, 80))
plt.colorbar(im2, ax=ax2, label='Strain (μm)')
ax2.set_xlabel('Time (s)'); ax2.set_ylabel('Depth (m)')
ax2.set_title(f'Strain Waterfall (Bandpass {bp_low:.4f}–{bp_high:.4f} Hz) — {ch_avg}-ch section windows shown')
draw_section_lines(ax2, orientation='depth')
ax2.legend(fontsize=7, loc='upper right')
plt.tight_layout()
savefig(fig2, figures_dir, 'fig2_strain_waterfall_filtered.png')

del wf_data, wf_filt

# ============================================================
# FIGURE 3 — FFT Amplitude Spectrum (one subplot per section, dynamic)
# ============================================================
fmax_show = min(5 * Freq, fs_d / 2)
fmask     = f_axis <= fmax_show
n_sections = len(all_channels)

fig3, axes3 = plt.subplots(n_sections, 1,
                            figsize=(9, 3 * n_sections),
                            sharex=True)
if n_sections == 1:
    axes3 = [axes3]

for ax, ch in zip(axes3, all_channels):
    half  = ch_avg // 2
    ch_lo = max(0, ch - half)
    ch_hi = min(nch, ch + half)
    rate  = subdata[ch_lo:ch_hi, :].mean(axis=0).astype(np.float64)
    rate -= rate.mean()
    disp   = cumtrapz(rate, dx=1/fs_d) * 1e-3
    d_disp = detrend(disp, type='linear')
    Amp    = 2 * np.abs(np.fft.fft(d_disp, n=NFFT)[:NFFT//2+1]) / NFFT
    ax.plot(f_axis[fmask], Amp[fmask], lw=1.5, color=section_color(ch))
    ax.axvline(Freq, color='k', linestyle='--', alpha=0.5, label=f'{Freq:.4f} Hz')
    ax.set_ylabel('Amplitude (μm)')
    ax.set_title(f'{section_label(ch)} — {ch_avg}-ch avg, unfiltered disp')
    ax.legend(fontsize=7); ax.grid(True)

axes3[-1].set_xlabel('Frequency (Hz)')
axes3[-1].set_xlim([0, fmax_show])
fig3.suptitle(f'Amplitude Spectra — All Sections ({ch_avg}-ch average)', fontsize=12)
plt.tight_layout()
savefig(fig3, figures_dir, 'fig3_fft_spectrum.png')

# ============================================================
# FIGURE 3B — FFT of Strain Rate (one subplot per section)
# ============================================================
fig3b, axes3b = plt.subplots(n_sections, 1,
                              figsize=(9, 3 * n_sections),
                              sharex=True)
if n_sections == 1:
    axes3b = [axes3b]

for ax, ch in zip(axes3b, all_channels):
    half  = ch_avg // 2
    ch_lo = max(0, ch - half)
    ch_hi = min(nch, ch + half)
    rate  = subdata[ch_lo:ch_hi, :].mean(axis=0).astype(np.float64)
    rate -= rate.mean()
    d_rate = detrend(rate, type='linear')
    Amp    = 2 * np.abs(np.fft.fft(d_rate, n=NFFT)[:NFFT//2+1]) / NFFT
    ax.plot(f_axis[fmask], Amp[fmask], lw=1.5, color=section_color(ch))
    ax.axvline(Freq, color='k', linestyle='--', alpha=0.5, label=f'{Freq:.4f} Hz')
    ax.set_ylabel('Amplitude (nm/s)')
    ax.set_title(f'{section_label(ch)} — {ch_avg}-ch avg, strain rate')
    ax.legend(fontsize=7); ax.grid(True)

axes3b[-1].set_xlabel('Frequency (Hz)')
axes3b[-1].set_xlim([0, fmax_show])
fig3b.suptitle(f'Strain Rate Spectra — All Sections ({ch_avg}-ch average)', fontsize=12)
plt.tight_layout()
savefig(fig3b, figures_dir, 'fig3b_strain_rate_spectrum.png')
# ============================================================
# PER-CHANNEL HELPER — averages ch_avg channels centered on ch_idx
# ============================================================
def channel_pp_amplitudes(ch_idx):
    half    = ch_avg // 2
    ch_lo   = max(0, ch_idx - half)
    ch_hi   = min(nch, ch_idx + half)
    rate    = subdata[ch_lo:ch_hi, :].mean(axis=0).astype(np.float64)
    rate   -= rate.mean()
    disp      = cumtrapz(rate, dx=1/fs_d) * 1e-3
    disp_filt = bandpass(disp, fs_d, bp_low, bp_high, bp_order)
    d_disp    = detrend(disp, type='linear')
    _, _, amp_raw = peak_envelope(d_disp,    fs_d, Freq)
    _, _, amp_bp  = peak_envelope(disp_filt, fs_d, Freq)
    return disp, disp_filt, d_disp, float(np.nanmean(amp_raw)), float(np.nanmean(amp_bp))

# ============================================================
# FIGURE 4 — Displacement per cable section
# ============================================================
DispOut, DispOut_filt, dDispOut, aOut_raw, aOut_bp = channel_pp_amplitudes(ref_channel)

section_disps = {}
for ch in inner_channels:
    disp, disp_filt, d_disp, a_raw, a_bp = channel_pp_amplitudes(ch)
    section_disps[ch] = {
        'disp': disp, 'disp_filt': disp_filt, 'd_disp': d_disp,
        'a_raw': a_raw, 'a_bp': a_bp,
        'ratio_raw': a_raw / aOut_raw if aOut_raw > 0 else np.nan,
        'ratio_bp':  a_bp  / aOut_bp  if aOut_bp  > 0 else np.nan,
    }

fig4, axes = plt.subplots(2, 1, figsize=(10, 7), sharex=True)
axes[0].plot(t_plot, DispOut[::plot_step],
             color=section_color(ref_channel), lw=1.5, label=section_label(ref_channel))
for ch in inner_channels:
    axes[0].plot(t_plot, section_disps[ch]['disp'][::plot_step],
                 color=section_color(ch), lw=1.2, label=section_label(ch))
axes[0].set_ylabel('Displacement (μm)')
axes[0].set_title(f'Mean Displacement — Unfiltered ({ch_avg}-ch average per section)')
axes[0].legend(fontsize=8); axes[0].grid(True)

axes[1].plot(t_plot, DispOut_filt[::plot_step],
             color=section_color(ref_channel), lw=1.5, label=section_label(ref_channel))
for ch in inner_channels:
    axes[1].plot(t_plot, section_disps[ch]['disp_filt'][::plot_step],
                 color=section_color(ch), lw=1.2, label=section_label(ch))
axes[1].set_ylabel('Displacement (μm)'); axes[1].set_xlabel('Time (s)')
axes[1].set_title(f'Mean Displacement — Bandpass {bp_low:.4f}–{bp_high:.4f} Hz ({ch_avg}-ch average per section)')
axes[1].legend(fontsize=8); axes[1].grid(True)
plt.tight_layout()
savefig(fig4, figures_dir, 'fig4_mean_displacement.png')

# ============================================================
# FIGURE 5 — Envelopes: each inner cable vs reference (dynamic rows)
# ============================================================
n_inner = len(inner_channels)
fig5, axes5 = plt.subplots(n_inner, 2, figsize=(12, 3.5 * n_inner), sharex='col')
if n_inner == 1:
    axes5 = axes5[np.newaxis, :]

for row, ch in enumerate(inner_channels):
    sd = section_disps[ch]
    up_out,    lo_out,    _ = peak_envelope(dDispOut,        fs_d, Freq)
    up_in,     lo_in,     _ = peak_envelope(sd['d_disp'],    fs_d, Freq)
    up_out_bp, lo_out_bp, _ = peak_envelope(DispOut_filt,    fs_d, Freq)
    up_in_bp,  lo_in_bp,  _ = peak_envelope(sd['disp_filt'], fs_d, Freq)

    col_ref = section_color(ref_channel)
    col_in  = section_color(ch)

    ax_raw = axes5[row, 0]
    ax_raw.plot(t_plot, dDispOut[::plot_step],         color=col_ref, label=section_label(ref_channel))
    ax_raw.plot(t_plot, up_out[::plot_step],           color=col_ref, lw=0.8, linestyle='--', alpha=0.7)
    ax_raw.plot(t_plot, lo_out[::plot_step],           color=col_ref, lw=0.8, linestyle='--', alpha=0.7)
    ax_raw.plot(t_plot, sd['d_disp'][::plot_step],     color=col_in,  label=section_label(ch))
    ax_raw.plot(t_plot, up_in[::plot_step],            color=col_in,  lw=0.8, linestyle='--', alpha=0.7)
    ax_raw.plot(t_plot, lo_in[::plot_step],            color=col_in,  lw=0.8, linestyle='--', alpha=0.7)
    ax_raw.set_ylabel('Disp (μm)')
    ax_raw.set_title(f'Unfilt {section_name(ch)} vs Ref: {sd["ratio_raw"]:.4f} ({sd["ratio_raw"]*100:.1f}%)')
    ax_raw.legend(fontsize=7); ax_raw.grid(True)

    ax_bp = axes5[row, 1]
    ax_bp.plot(t_plot, DispOut_filt[::plot_step],        color=col_ref, label=section_label(ref_channel))
    ax_bp.plot(t_plot, up_out_bp[::plot_step],           color=col_ref, lw=0.8, linestyle='--', alpha=0.7)
    ax_bp.plot(t_plot, lo_out_bp[::plot_step],           color=col_ref, lw=0.8, linestyle='--', alpha=0.7)
    ax_bp.plot(t_plot, sd['disp_filt'][::plot_step],     color=col_in,  label=section_label(ch))
    ax_bp.plot(t_plot, up_in_bp[::plot_step],            color=col_in,  lw=0.8, linestyle='--', alpha=0.7)
    ax_bp.plot(t_plot, lo_in_bp[::plot_step],            color=col_in,  lw=0.8, linestyle='--', alpha=0.7)
    ax_bp.set_ylabel('Disp (μm)')
    ax_bp.set_title(f'BP {section_name(ch)} vs Ref: {sd["ratio_bp"]:.4f} ({sd["ratio_bp"]*100:.1f}%)')
    ax_bp.legend(fontsize=7); ax_bp.grid(True)

axes5[-1, 0].set_xlabel('Time (s)')
axes5[-1, 1].set_xlabel('Time (s)')
fig5.suptitle(f'Strain Transfer Ratios vs Reference ({ch_avg}-ch average per section)', fontsize=12)
plt.tight_layout()
savefig(fig5, figures_dir, 'fig5_envelopes.png')

# ============================================================
# FIGURE 6 — Channel amplitude profile along fiber
# ============================================================
print(f"  Computing Fig 6 per-channel amplitudes (ds={fig6_ds})...")
fs_ds = fs_d / fig6_ds
ch_list, amp_raw_list, amp_bp_list = [], [], []

for ch in range(nch):
    rate = subdata[ch, ::fig6_ds].astype(np.float64)
    rate -= rate.mean()
    disp         = cumtrapz(rate, dx=fig6_ds/fs_d) * 1e-3
    disp_filt    = bandpass(disp, fs_ds, bp_low, bp_high, bp_order)
    disp_detrend = detrend(disp, type='linear')
    _, _, amp_r  = peak_envelope(disp_detrend, fs_ds, Freq)
    _, _, amp_b  = peak_envelope(disp_filt,    fs_ds, Freq)
    ch_list.append(ch)
    amp_raw_list.append(float(np.nanmean(amp_r)))
    amp_bp_list.append(float(np.nanmean(amp_b)))

ch_arr      = np.array(ch_list)
amp_raw_arr = np.array(amp_raw_list)
amp_bp_arr  = np.array(amp_bp_list)

fig6, ax6 = plt.subplots(figsize=(14, 5))
ax6.plot(ch_arr, amp_raw_arr, 'b-', lw=1.5, label='Unfiltered')
ax6.plot(ch_arr, amp_bp_arr,  'r-', lw=1.5, label=f'Bandpass {bp_low:.4f}–{bp_high:.4f} Hz')
draw_section_lines(ax6, orientation='channel')
ax6.set_xlabel('Channel')
ax6.set_ylabel('Displacement Amplitude (μm)')
ax6.set_title(f'Channel Amplitude Profile — Integration Method ({ch_avg}-ch section windows shown)')
ax6.xaxis.set_major_locator(MultipleLocator(100))
ax6.xaxis.set_minor_locator(MultipleLocator(10))
ax6.legend(fontsize=7, loc='upper right')
ax6.grid(True)
plt.tight_layout()
savefig(fig6, figures_dir, 'fig6_channel_amplitude_profile.png')

# ============================================================
# SUMMARY
# ============================================================
print("\n========== STRAIN TRANSFER SUMMARY ==========")
print(f"  File                : {load_label}")
print(f"  Figures saved to    : {figures_dir}")
print(f"  Signal frequency    : {Freq:.4f} Hz")
print(f"  Bandpass            : {bp_low:.4f} – {bp_high:.4f} Hz")
print(f"  Channel averaging   : {ch_avg} channels per section")
print(f"  Depth crop          : {max_depth_m} m  ({ch_stop} channels)")
print(f"  Reference           : {section_label(ref_channel)}")
print(f"  Test sections       : {[section_label(ch) for ch in inner_channels]}")
print(f"  Total channels      : {n_channels}")
print(f"  P-P Reference (unfilt / BP): {aOut_raw:.3f} / {aOut_bp:.3f} μm")
print(f"  {'Section':<12} {'P-P unfilt':>12} {'P-P BP':>12} {'Ratio unfilt':>14} {'Ratio BP':>14}")
print(f"  {'-'*66}")
for ch in inner_channels:
    sd = section_disps[ch]
    print(
        f"  {section_name(ch):<12} {sd['a_raw']:>10.3f} μm {sd['a_bp']:>10.3f} μm"
        f" {sd['ratio_raw']:>12.4f} {sd['ratio_bp']:>12.4f}"
    )
print("==============================================\n")

# ============================================================
# UPDATE EXCEL RESULTS FILE
# ============================================================
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if freq_val is None or voltage_val is None:
    print(f"  WARNING: Could not parse freq/voltage from path ({freq_folder}/{vol_folder})")
    print(f"  Excel not updated. Check folders are named like '1hz' and '7V'.")
else:
    print(f"  Logging to Excel: {freq_val} Hz / {voltage_val} V")

    sheet_defs = {
        'Reference P-P Unfiltered': aOut_raw,
        'Reference P-P Bandpass':   aOut_bp,
    }
    for ch in inner_channels:
        name = section_name(ch)
        sd   = section_disps[ch]
        sheet_defs[f'{name} P-P Unfiltered']   = sd['a_raw']
        sheet_defs[f'{name} P-P Bandpass']     = sd['a_bp']
        sheet_defs[f'{name} Ratio Unfiltered'] = sd['ratio_raw']
        sheet_defs[f'{name} Ratio Bandpass']   = sd['ratio_bp']

    sheet_order = ['Reference P-P Unfiltered', 'Reference P-P Bandpass']
    for ch in inner_channels:
        name = section_name(ch)
        sheet_order.extend([
            f'{name} P-P Unfiltered', f'{name} P-P Bandpass',
            f'{name} Ratio Unfiltered', f'{name} Ratio Bandpass',
        ])

    header_fill  = PatternFill('solid', start_color='1F4E79')
    header_font  = Font(bold=True, color='FFFFFF', name='Arial', size=11)
    index_fill   = PatternFill('solid', start_color='D6E4F0')
    index_font   = Font(bold=True, name='Arial', size=10)
    body_font    = Font(name='Arial', size=10)
    center_align = Alignment(horizontal='center', vertical='center')
    thin         = Side(style='thin')
    bdr          = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_cell(cell, font=None, fill=None):
        cell.alignment = center_align
        cell.border    = bdr
        if font: cell.font = font
        if fill: cell.fill = fill

    def write_sorted_grid(ws, freq_val, voltage_val, value):
        grid = {}
        voltages, freqs = set(), set()
        for col in range(2, max(ws.max_column, 2) + 1):
            v = ws.cell(row=2, column=col).value
            if v is not None:
                voltages.add(float(v))
        for row in range(3, max(ws.max_row, 3) + 1):
            f = ws.cell(row=row, column=1).value
            if f is not None:
                freqs.add(float(f))
                grid[float(f)] = {}
                for col in range(2, ws.max_column + 1):
                    v = ws.cell(row=2, column=col).value
                    if v is not None:
                        grid[float(f)][float(v)] = ws.cell(row=row, column=col).value

        fv, vv = float(freq_val), float(voltage_val)
        grid.setdefault(fv, {})[vv] = round(value, 4) if not np.isnan(value) else 'N/A'
        voltages.add(vv); freqs.add(fv)

        all_volts = sorted(voltages)
        all_freqs = sorted(freqs)

        for row in range(3, ws.max_row + 1):
            for col in range(2, ws.max_column + 1):
                ws.cell(row=row, column=col).value = None
        for row in range(3, ws.max_row + 1):
            ws.cell(row=row, column=1).value = None

        style_cell(ws['A2'], font=header_font, fill=header_fill)
        for j, v in enumerate(all_volts, start=2):
            style_cell(ws.cell(row=2, column=j), font=header_font, fill=header_fill)
            ws.cell(row=2, column=j).value = v
        for i, f in enumerate(all_freqs, start=3):
            style_cell(ws.cell(row=i, column=1), font=index_font, fill=index_fill)
            ws.cell(row=i, column=1).value = f
            for j, v in enumerate(all_volts, start=2):
                cell_val = grid.get(f, {}).get(v)
                if cell_val is not None:
                    style_cell(ws.cell(row=i, column=j), font=body_font)
                    ws.cell(row=i, column=j).value = cell_val

        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 14)
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row == 1 and merged.max_row == 1:
                try:
                    ws.unmerge_cells(str(merged))
                except KeyError:
                    pass
        if len(all_volts) >= 1:
            ws.merge_cells(start_row=1, start_column=1,
                           end_row=1, end_column=len(all_volts) + 1)
            ws['A1'].alignment = center_align

    wb = load_workbook(excel_path) if os.path.exists(excel_path) else Workbook()
    if 'Sheet' in wb.sheetnames and len(wb.sheetnames) == 1:
        del wb['Sheet']

    for sheet_name, value in sheet_defs.items():
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            ws['A1'] = sheet_name
            ws['A1'].font      = Font(bold=True, name='Arial', size=13, color='1F4E79')
            ws['A1'].alignment = center_align
            ws['A2'] = 'Freq (Hz) \\ Voltage (V)'
        else:
            ws = wb[sheet_name]
        write_sorted_grid(ws, freq_val, voltage_val, value)

    for old_name in list(wb.sheetnames):
        if old_name not in sheet_defs:
            del wb[old_name]

    name_to_ws = {ws.title: ws for ws in wb.worksheets}
    wb._sheets = [name_to_ws[n] for n in sheet_order if n in name_to_ws]

    wb.save(excel_path)
    print(f"  Excel updated: {excel_path}")